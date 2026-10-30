"""Generate a small, synthetic all-formats demo pack for the Care Guide UI."""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "test sample"


def invoice(number: str | None, *, subtotal: str = "100.00", tax: str | None = "10.00",
            total: str = "110.00", customer: str = "Care Guide", supplier: str = "North Clinic",
            currency: str = "USD", items: list[dict[str, str]] | None = None) -> dict[str, Any]:
    return {
        "invoice_number": number,
        "invoice_date": "2026-05-10",
        "supplier_name": supplier,
        "customer_name": customer,
        "subtotal": subtotal,
        "tax_amount": tax,
        "total_amount": total,
        "currency": currency,
        "line_items": items or [
            {"description": "Care visit", "quantity": "2", "unit_price": "50.00", "line_total": "100.00"},
            {"description": "Diagnostic kit", "quantity": "1", "unit_price": "0.00", "line_total": "0.00"},
        ],
    }


VALID = invoice("CG-DEMO-1001")
BAD_TOTAL = invoice("CG-DEMO-1002", subtotal="120.00", tax="12.00", total="139.00",
                    items=[{"description": "Lab panel", "quantity": "2", "unit_price": "60.00", "line_total": "120.00"}])
BAD_LINES = invoice("CG-DEMO-1003", items=[
    {"description": "Care service", "quantity": "2", "unit_price": "45.00", "line_total": "100.00"},
    {"description": "Equipment", "quantity": "1", "unit_price": "0.00", "line_total": "5.00"},
])
SCANNED = invoice("CG-DEMO-1004", items=[
    {"description": "Scanned consultation", "quantity": "2", "unit_price": "50.00", "line_total": "100.00"},
])
MULTIPAGE = invoice("CG-DEMO-1004", items=[
    {"description": "Scanned consultation", "quantity": "2", "unit_price": "50.00", "line_total": "100.00"},
    {"description": "Diagnostic kit", "quantity": "1", "unit_price": "0.00", "line_total": "0.00"},
])


def _invoice_lines(data: dict[str, Any]) -> list[str]:
    lines = [
        "INVOICE",
        f"Invoice Number: {data['invoice_number']}" if data["invoice_number"] else "Invoice Number:",
        f"Invoice Date: {data['invoice_date']}",
        f"Supplier: {data['supplier_name']}",
        f"Customer: {data['customer_name']}",
        f"Currency: {data['currency']}",
        f"Subtotal: {data['subtotal']}",
    ]
    if data["tax_amount"] is not None:
        lines.append(f"Tax: {data['tax_amount']}")
    lines.extend([f"Total: {data['total_amount']}", "Description | Quantity | Unit Price | Line Total"])
    lines.extend(" | ".join((item["description"], item["quantity"], item["unit_price"], item["line_total"]))
                 for item in data["line_items"])
    return lines


def _write_pdf(path: Path, pages: list[list[str]], *, image: Image.Image | None = None) -> None:
    import fitz

    doc = fitz.open()
    for page_lines in pages:
        page = doc.new_page(width=612, height=792)
        if image is not None:
            page.insert_image(page.rect, stream=_image_bytes(image, "PNG"))
        else:
            for idx, line in enumerate(page_lines):
                page.insert_text((54, 62 + idx * 30), line, fontsize=13, fontname="helv")
    doc.set_metadata({"title": "Synthetic Care Guide invoice", "author": "Care Guide demo fixtures"})
    doc.save(path, garbage=4, deflate=True)
    doc.close()


def _image_bytes(image: Image.Image, image_format: str) -> bytes:
    from io import BytesIO

    out = BytesIO()
    image.save(out, format=image_format)
    return out.getvalue()


def _font(size: int) -> ImageFont.ImageFont:
    font_file = Path("C:/Windows/Fonts/arial.ttf")
    return ImageFont.truetype(str(font_file), size) if font_file.is_file() else ImageFont.load_default()


def _make_invoice_image(data: dict[str, Any], *, blank_number: bool = False,
                        page_label: str = "") -> Image.Image:
    image = Image.new("RGB", (1650, 2200), "white")
    draw = ImageDraw.Draw(image)
    font = _font(46)
    small = _font(38)
    y = 90
    rows = [
        ("NORTH CLINIC", font), ("INVOICE", font),
        (f"Invoice Number: {' ' if blank_number else data['invoice_number']}", small),
        (f"Invoice Date: {data['invoice_date']}", small),
        (f"Supplier: {data['supplier_name']}", small),
        (f"Customer: {data['customer_name']}", small),
        (f"Currency: {data['currency']}", small),
        (f"Subtotal: {data['subtotal']}", small),
    ]
    if data["tax_amount"] is not None:
        rows.append((f"Tax: {data['tax_amount']}", small))
    rows.extend([(f"Total: {data['total_amount']}", font),
                 ("Description   Quantity   Unit Price   Line Total", small)])
    rows.extend((f"{item['description']}   {item['quantity']}   {item['unit_price']}   {item['line_total']}", small)
                for item in data["line_items"])
    if page_label:
        rows.append((page_label, small))
    for text, selected_font in rows:
        draw.text((100, y), text, fill="#111111", font=selected_font)
        y += 130
    return image


def _write_csv(path: Path, data: dict[str, Any], *, note: str | None = None) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        for key, value in [
            ("Invoice Number", data["invoice_number"]), ("Invoice Date", data["invoice_date"]),
            ("Supplier", data["supplier_name"]), ("Customer", data["customer_name"]),
            ("Currency", data["currency"]), ("Subtotal", data["subtotal"]),
            ("Tax", data["tax_amount"]), ("Total", data["total_amount"]),
        ]:
            if value is not None:
                writer.writerow([key, value])
        if note:
            writer.writerow(["Export note", note])
        writer.writerow([])
        writer.writerow(["Description", "Quantity", "Unit Price", "Line Total"])
        for item in data["line_items"]:
            writer.writerow([item["description"], item["quantity"], item["unit_price"], item["line_total"]])


def _write_xlsx(path: Path, data: dict[str, Any], *, line_items: list[dict[str, str]] | None = None) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    header = wb.active
    header.title = "Invoice"
    for key, value in [
        ("Invoice Number", data["invoice_number"]), ("Invoice Date", data["invoice_date"]),
        ("Supplier", data["supplier_name"]), ("Customer", data["customer_name"]),
        ("Currency", data["currency"]), ("Subtotal", data["subtotal"]),
        ("Tax", data["tax_amount"]), ("Total", data["total_amount"]),
    ]:
        if value is not None:
            header.append([key, value])
    items_sheet = wb.create_sheet("Line Items")
    items_sheet.append(["Description", "Quantity", "Unit Price", "Line Total"])
    for item in line_items or data["line_items"]:
        items_sheet.append([item["description"], item["quantity"], item["unit_price"], item["line_total"]])
    wb.save(path)


def _write_xls(path: Path, data: dict[str, Any]) -> None:
    import xlwt

    wb = xlwt.Workbook()
    header = wb.add_sheet("Invoice")
    for row, (key, value) in enumerate([
        ("Invoice Number", data["invoice_number"]), ("Invoice Date", data["invoice_date"]),
        ("Supplier", data["supplier_name"]), ("Customer", data["customer_name"]),
        ("Currency", data["currency"]), ("Subtotal", data["subtotal"]),
        ("Tax", data["tax_amount"]), ("Total", data["total_amount"]),
    ]):
        if value is not None:
            header.write(row, 0, key)
            header.write(row, 1, value)
    items = wb.add_sheet("Line Items")
    for col, label in enumerate(["Description", "Quantity", "Unit Price", "Line Total"]):
        items.write(0, col, label)
    for row, item in enumerate(data["line_items"], 1):
        for col, key in enumerate(["description", "quantity", "unit_price", "line_total"]):
            items.write(row, col, item[key])
    wb.save(str(path))


def _write_long_pdf(path: Path, pages: int = 1001) -> None:
    import fitz

    doc = fitz.open()
    for page_number in range(1, pages + 1):
        page = doc.new_page(width=612, height=792)
        page.insert_text((54, 72), f"Synthetic demo page {page_number} of {pages}", fontsize=12)
    doc.set_metadata({"title": "Optional 1001 page page-limit fixture"})
    doc.save(path, garbage=4, deflate=True)
    doc.close()


def _manifest() -> dict[str, Any]:
    mismatch = invoice("CG-DEMO-1001", total="111.00")
    missing_number = invoice(None)
    missing_tax = invoice("CG-DEMO-1005", tax=None, total="100.00")
    uncertain = invoice(None)
    return {
        "dataset_name": "Care Guide local UI demo samples",
        "synthetic_only": True,
        "normalization": "Dates are ISO YYYY-MM-DD; money values are exact decimal strings.",
        "comparison_pairs": {
            "matching": ["01_valid_invoice.pdf", "02_matching_invoice.xlsx"],
            "mismatch": ["01_valid_invoice.pdf", "03_mismatch_same_invoice.csv"],
        },
        "documents": [
            {"file": "01_valid_invoice.pdf", "expected_invoice": VALID, "expected_validation_codes": []},
            {"file": "02_matching_invoice.xlsx", "expected_invoice": VALID, "expected_validation_codes": ["duplicate_invoice"]},
            {"file": "03_mismatch_same_invoice.csv", "expected_invoice": mismatch, "expected_validation_codes": ["duplicate_invoice", "subtotal_tax_total_mismatch"]},
            {"file": "04_duplicate_invoice.csv", "expected_invoice": VALID, "expected_validation_codes": ["duplicate_invoice"]},
            {"file": "05_missing_invoice_number.csv", "expected_invoice": missing_number, "expected_validation_codes": ["required_field_missing"]},
            {"file": "06_missing_tax.csv", "expected_invoice": missing_tax, "expected_validation_codes": ["field_missing"]},
            {"file": "07_incorrect_total.pdf", "expected_invoice": BAD_TOTAL, "expected_validation_codes": ["subtotal_tax_total_mismatch"]},
            {"file": "08_line_item_error.xlsx", "expected_invoice": BAD_LINES, "expected_validation_codes": ["line_item_multiplication_mismatch", "line_items_subtotal_mismatch"]},
            {"file": "09_uncertain_value.png", "expected_invoice": uncertain, "unreadable_fields": ["invoice_number"], "requires_ocr": True},
            {"file": "10_scanned_invoice.jpg", "expected_invoice": SCANNED, "requires_ocr": True},
            {"file": "11_scanned_invoice.pdf", "expected_invoice": SCANNED, "requires_ocr": True},
            {"file": "12_scanned_invoice.png", "expected_invoice": SCANNED, "requires_ocr": True},
            {"file": "13_multipage_invoice.tiff", "expected_invoice": MULTIPAGE, "page_count": 2, "requires_ocr": True},
            {"file": "14_legacy_invoice.xls", "expected_invoice": VALID, "expected_worksheets": ["Invoice", "Line Items"]},
            {"file": "15_malformed_rows.csv", "expected_processing_status": "completed_with_errors"},
            {"file": "16_empty.csv", "expected_upload_status": "rejected", "expected_error": "empty_file"},
            {"file": "17_corrupted.pdf", "expected_processing_status": "failed"},
            {"file": "18_unsupported.txt", "expected_upload_status": "rejected", "expected_error": "unsupported_format"},
            {"file": "optional/optional_1001_page_limit.pdf", "page_count": 1001, "expected_processed_pages": 1000, "expected_warning": "page_limit_reached"},
        ],
        "retry_note": "A source fixture cannot reliably produce a transient worker failure. The corrupt PDF is permanently invalid and should not be retryable. Use the automated worker tests for transient retry behavior.",
        "optional_features": "Ollama suggestions and semantic search need optional local models; their absence does not prevent ordinary upload, extraction, review, keyword search, or comparisons.",
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "optional").mkdir(exist_ok=True)
    (OUT / "01_valid_invoice.pdf").write_bytes(b"")
    _write_pdf(OUT / "01_valid_invoice.pdf", [
        _invoice_lines(VALID)[:6],
        _invoice_lines(VALID)[6:],
    ])
    _write_xlsx(OUT / "02_matching_invoice.xlsx", VALID)
    _write_csv(OUT / "03_mismatch_same_invoice.csv", invoice("CG-DEMO-1001", total="111.00"), note="Different source export")
    _write_csv(OUT / "04_duplicate_invoice.csv", VALID, note="Duplicate export copy, not identical bytes")
    _write_csv(OUT / "05_missing_invoice_number.csv", invoice(None))
    _write_csv(OUT / "06_missing_tax.csv", invoice("CG-DEMO-1005", tax=None, total="100.00"))
    _write_pdf(OUT / "07_incorrect_total.pdf", [_invoice_lines(BAD_TOTAL)])
    _write_xlsx(OUT / "08_line_item_error.xlsx", BAD_LINES)

    uncertain_img = _make_invoice_image(invoice(None), blank_number=True)
    uncertain_img.save(OUT / "09_uncertain_value.png", format="PNG", optimize=True)
    scanned_img = _make_invoice_image(SCANNED)
    scanned_img.save(OUT / "10_scanned_invoice.jpg", format="JPEG", quality=94, optimize=True)
    scanned_img.save(OUT / "12_scanned_invoice.png", format="PNG", optimize=True)
    _write_pdf(OUT / "11_scanned_invoice.pdf", [[]], image=scanned_img)
    first = _make_invoice_image(SCANNED, page_label="Page 1 of 2")
    second_data = invoice("CG-DEMO-1004", items=[
        {"description": "Diagnostic kit", "quantity": "1", "unit_price": "0.00", "line_total": "0.00"},
    ])
    second = _make_invoice_image(second_data, page_label="Page 2 of 2")
    first.save(OUT / "13_multipage_invoice.tiff", format="TIFF", save_all=True, append_images=[second], compression="tiff_lzw")

    _write_xls(OUT / "14_legacy_invoice.xls", VALID)
    with (OUT / "15_malformed_rows.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["Invoice Number", "CG-DEMO-MALFORMED"])
        writer.writerow([])
        writer.writerow(["Description", "Quantity", "Unit Price", "Line Total"])
        writer.writerow(["Care visit", "2", "50.00", "100.00"])
        writer.writerow(["A malformed extra column", "1", "0.00", "0.00", "unexpected"])
    (OUT / "16_empty.csv").write_bytes(b"")
    (OUT / "17_corrupted.pdf").write_bytes(b"%PDF-1.7\nSynthetic corrupt fixture; no valid PDF objects.\n")
    (OUT / "18_unsupported.txt").write_text("Harmless unsupported test file.\n", encoding="utf-8")
    _write_long_pdf(OUT / "optional" / "optional_1001_page_limit.pdf")

    with (OUT / "ground_truth.json").open("w", encoding="utf-8", newline="\n") as stream:
        json.dump(_manifest(), stream, indent=2, ensure_ascii=False)
        stream.write("\n")
    (OUT / "README.txt").write_text(README, encoding="utf-8")
    print(f"Generated synthetic test samples in: {OUT}")


README = """CARE GUIDE SYNTHETIC TEST SAMPLES
===================================

These files are fictional test data. Upload the files to the Care Guide app; do not treat them as real invoices.

Suggested end-to-end demo:
1. Upload 01_valid_invoice.pdf, 02_matching_invoice.xlsx, 05_missing_invoice_number.csv,
   06_missing_tax.csv, 07_incorrect_total.pdf, and 11_scanned_invoice.pdf together.
2. Keep the API and worker running until the rows show completed/failed. Open each row and inspect fields,
   line items, review findings, processing history, and (for the scan) OCR output.
3. Compare 01_valid_invoice.pdf with 02_matching_invoice.xlsx; then compare 01_valid_invoice.pdf
   with 03_mismatch_same_invoice.csv. The first pair should match; the second shares invoice identity
   but has a different total.
4. Search for North Clinic with Keyword mode.
5. Open 07_incorrect_total.pdf, correct Total amount to 132.00 (120.00 subtotal + 12.00 tax), enter a
   reason, and choose Save review. Inspect More details and download the reviewed result.
6. Upload 16_empty.csv, 17_corrupted.pdf, and 18_unsupported.txt to see isolated rejection/failure handling.
7. Upload 14_legacy_invoice.xls and 15_malformed_rows.csv to inspect legacy workbook and malformed CSV handling.
8. Optional: upload optional/optional_1001_page_limit.pdf to demonstrate the configured 1,000-page limit.

Files 01 and 02 are a same-invoice comparison pair. Files 01 and 03 are a same-identity pair with a
deliberate total mismatch. Files 02, 03, and 04 also exercise duplicate invoice detection; the app should
not decide that documents are duplicates based only on equal totals.

Retry behavior: the corrupt PDF is permanently invalid, so it is correctly not retryable. A sample file
cannot reliably create a transient worker error. Use the automated worker retry tests for that state, or
upload a valid file while the worker is stopped and restart the worker to show queued work resumes.

Optional runtime AI: Ollama suggestions require Ollama and a local model. Semantic search requires the
optional sentence-transformers package and model weights. Ordinary deterministic extraction and keyword
search work without either. Do not select semantic mode in a demo unless the model is installed and ready.

Ground-truth expectations are in ground_truth.json. OCR can make character-level mistakes, so review the
extracted values rather than assuming every scanned field will be exact.

To regenerate from the project root, run:
    .venv\\Scripts\\python.exe scripts\\generate_test_samples.py

The optional 1,001-page file is separate so it does not slow down the ordinary multi-file demo.
"""


if __name__ == "__main__":
    main()
